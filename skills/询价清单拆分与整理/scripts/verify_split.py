# -*- coding: utf-8 -*-
"""独立校验：回读源 .xls 与输出 .xlsx，按表头名定位列，做「身份级」比对。

用法:  python verify_split.py [strict|loose] [multi|single]

设计要点：**校验侧独立采用语义正确的分组键**，不复用生成脚本的键。
关键差异（正是这处差异抓出过一次真实误合并）：
  - 分组键必须包含**技术参数全文**。散热器 R39/R40/R41（15/20/25 片/组）、
    排风机 R67/R68（两种工况）的差异全在技术参数列，只按「规格型号」分组会把它们误合。
  - 源侧用**原始规格型号整串**归一，输出侧用「规格型号+技术参数」拼接后归一，两者必须相等。

校验项：
 1) 源 273 行按模式口径分组后的 (分组, 材质, 数量) 多重集 == 输出；
 2) 数量总和守恒；3) 无遗漏；4) 无捏造（代表值必须来自成员原文）；
 5) 格式面：表头/冻结/筛选/合计公式/序号连续/单价留空/合价公式/标题地址。
"""
import xlrd, re, sys, io, os, glob
from collections import Counter, defaultdict, OrderedDict
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

MODE = (sys.argv[1] if len(sys.argv) > 1 else "strict").lower().strip()
assert MODE in ("strict", "loose"), MODE
LAYOUT = (sys.argv[2] if len(sys.argv) > 2 else "multi").lower().strip()
assert LAYOUT in ("multi", "single"), LAYOUT

SRC = r"D:\UserData\Documents\xwechat_files\wxid_ofdozq3goua321_291d\msg\file\2026-10\锌粉、熔铸车间设备、材料非标制作询价表.xls"
BASE = r"D:\UserData\Documents\Workbuddy\2026-10-12-15-43-34\锌粉熔铸车间_分类拆分"
OUT = BASE + ("" if MODE == "strict" else "_宽松")
ADDR = "济源市莲东村北河南豫光锌业有限公司厂区内"
print(f"=== 校验口径: {MODE} | 形态: {LAYOUT} | 目录: {OUT} ===\n")

FAIL = []
def check(ok, msg):
    print(("  [OK]   " if ok else "  [FAIL] ") + msg)
    if not ok: FAIL.append(msg)

def T(x):
    if x is None: return ""
    if isinstance(x, bool): return str(x)
    if isinstance(x, float) and x.is_integer(): return str(int(x))
    if isinstance(x, int): return str(x)
    return str(x).strip()

_NZ = lambda s: re.sub(r"[\s；;，,\n]", "", T(s))          # 严格口径：只剔除被列边界取代的字符
_ATLAS = re.compile(r"^\d{2}[A-Za-z]\d{3}$")

def canon_txt(s):
    """宽松口径：去空白/分隔符、全角转半角、PN= 视同 PN。"""
    t = T(s).replace("＝", "=").replace("：", ":")
    t = re.sub(r"[\s，,、；;\n]", "", t)
    return re.sub(r"PN\s*=\s*", "PN", t, flags=re.I)

def canon_mat(m):
    t = T(m)
    return "" if _ATLAS.match(t) else t

def canon_note(note, mat):
    t = re.sub(r"[\s、，,；;]", "", T(note))
    if t and t == re.sub(r"\s+", "", T(mat)):
        return ""
    return t

def spec_full(s):
    return canon_txt(s) if MODE == "loose" else _NZ(s)

def ident(name, spec, mat, unit, note):
    """返回 (分组键, 材质键)。分组键与生成脚本无关，独立按语义书写。"""
    if MODE == "loose":
        return (T(name), spec_full(spec), canon_note(note, mat), T(unit)), canon_mat(mat)
    return (T(name), spec_full(spec), T(mat), T(unit), T(note)), T(mat)

def absorb(pairs):
    """材质吸附：组内至多一种非空材质 → 空材质/图集号并入；多种真实材质 → 按材质各自成组。"""
    bym = OrderedDict()
    for m, q in pairs:
        bym[m] = bym.get(m, 0) + q
    nonempty = [m for m in bym if m]
    if len(nonempty) <= 1:
        return {(nonempty[0] if nonempty else ""): sum(bym.values())}
    return dict(bym)

def qsum(vals):
    t = sum(vals)
    return int(t) if float(t).is_integer() else round(t, 4)

# ---------- 排序键（独立实现；NAME_ORDER 只作数据从生成脚本读取）----------
_HERE = os.path.dirname(os.path.abspath(__file__))
# NAME_ORDER 在此仅作"数据"读取（排序规范）；生成脚本可能已改名为 split_by_category_example.py
_GEN = next((p for p in ("split_by_category_example.py", "gen_split.py")
             if os.path.exists(os.path.join(_HERE, p))), None)
assert _GEN, "未找到生成脚本（split_by_category_example.py / gen_split.py），无法读取 NAME_ORDER"
_src = open(os.path.join(_HERE, _GEN), encoding="utf-8").read()
NAME_ORDER = eval(re.search(r"NAME_ORDER = (\{(?:[^{}]|\{[^{}]*\})*\n\})", _src, re.S).group(1))
NAME_RANK = {(c, n): i for c, ns in NAME_ORDER.items() for i, n in enumerate(ns)}

_DNr = re.compile(r"DN\s*(\d+)", re.I)
_NUMr = re.compile(r"\d+(?:\.\d+)?")
_ANGLr = re.compile(r"^\s*(?:45|60|90|180)\s*(?:°|度)?\s*[A-Za-z]{0,3}\s*(?:\([A-Za-z]\))?")
_SCHr = re.compile(r"Sch\s*\d+", re.I)

def dim(spec):
    """口径键（独立重写，遵循与生成侧同一规范）：有 DN 取 DN 数字；
    否则剥掉角弯角度前缀与 Sch 壁厚系列后取数字。"""
    t = T(spec)
    d = _DNr.findall(t)
    if d:
        return tuple(int(x) for x in d)
    return tuple(float(x) for x in _NUMr.findall(_SCHr.sub("", _ANGLr.sub("", t))))

# ---------- 源侧 ----------
sh = xlrd.open_workbook(SRC).sheet_by_index(0)
def sv(r, c):
    v = sh.cell_value(r, c)
    return "" if v in (None, "") else (str(v).strip() if isinstance(v, str) else v)

src_coarse = defaultdict(list)      # 分组键 -> [(材质键, 数量), ...]
src_members = defaultdict(list)     # 分组键 -> [源明细, ...]
n_src = 0
for r in range(2, sh.nrows):
    name, spec, mat, unit, qty, note, extra = (sv(r, c) for c in range(1, 8))
    name, spec, mat, unit, note, extra = T(name), T(spec), T(mat), T(unit), T(note), T(extra)
    if not name:
        name, spec = spec, ""                       # 名称为空 → 系统名回填
    nf = (note + "；" + extra) if (note and extra) else (note or extra)
    k, mk = ident(name, spec, mat, unit, nf)
    q = float(qty); q = int(q) if q == int(q) else q
    src_coarse[k].append((mk, q))
    src_members[k].append(dict(er=r + 1, name=name, spec=spec, mat=mat, note=nf,
                              unit=unit, qty=q))
    n_src += 1

src_final = {}
for k, pairs in src_coarse.items():
    for m, q in absorb(pairs).items():
        src_final[(k, m)] = q

src_total = qsum([x["qty"] for ms in src_members.values() for x in ms])
print(f"源数据行 {n_src} | 分组键 {len(src_coarse)} | 合并后身份 {len(src_final)} | 数量总和 {src_total}")
print(f"重复组 {len(src_coarse) - len(src_final)} | 应消除 {n_src - len(src_final)} 行")

# ---------- 输出侧 ----------
HDR_TPL = ["序号", "类别", "名称", "规格型号", "材质/牌号", "技术参数",
           "单位", "数量", "备注", "单价", "合价"]
CAT_SET = {"钢材", "管件", "阀门", "桥架", "配电柜", "线缆", "其他", "设备", "仪表"}
CAT_ORDER = ["钢材", "管件", "阀门", "其他", "设备", "仪表"]     # 源表有数据的类别，出表顺序

# 收集待校验单元：(显示名, worksheet, 类别名, 是否分文件形态)
units = []
if LAYOUT == "multi":
    files = sorted(f for f in glob.glob(os.path.join(OUT, "*）.xlsx"))
                   if "全类别汇总" not in f)
    print(f"\n[形态 multi] 分文件 {len(files)} 个")
    check(len(files) == 6, f"分文件数 {len(files)} == 6（桥架/配电柜/线缆 源表 0 条不生成）")
    for fp in files:
        wb = load_workbook(fp)
        ws = wb[wb.sheetnames[0]]
        cat_name = os.path.basename(fp).split("（")[-1].rstrip("）.xlsx")
        units.append((os.path.basename(fp), ws, cat_name, True))
else:
    hit = glob.glob(os.path.join(OUT, "*（全类别汇总）.xlsx"))
    print(f"\n[形态 single] 合并文件 {len(hit)} 个")
    check(len(hit) == 1, f"合并文件存在且唯一（{len(hit)} 个）")
    fp = hit[0]
    wb = load_workbook(fp)
    print(f"  工作表 {len(wb.sheetnames)} 张: {wb.sheetnames}")
    check(wb.sheetnames == CAT_ORDER, f"工作表及顺序 == {CAT_ORDER}（实际 {wb.sheetnames}）")
    for nm in wb.sheetnames:
        units.append((os.path.basename(fp), wb[nm], wb[nm].title, False))

out_coarse = defaultdict(list)
out_members = defaultdict(list)
n_out = 0

for fname, ws, cat_name, is_multi in units:
    print(f"\n=== {fname} | sheet='{ws.title}' | {ws.dimensions} | 冻结={ws.freeze_panes} | 筛选={ws.auto_filter.ref}")
    if is_multi:
        check(ws.title == cat_name, f"{ws.title}: 工作表名 == 文件名类别")
    else:
        check(cat_name in CAT_SET, f"{ws.title}: 工作表名为已知类别")
    check([ws.cell(2, i).value for i in range(1, 12)] == HDR_TPL, f"{ws.title}: 表头 11 列与模板一致")
    check("A1:K1" in [str(m) for m in ws.merged_cells.ranges], f"{ws.title}: A1:K1 标题已合并")
    title = str(ws.cell(1, 1).value or "")
    check(ADDR in title and "＿＿" not in title and "___" not in title,
          f"{ws.title}: 标题含厂区地址且无占位符")

    last = ws.max_row
    check(ws.cell(last, 1).value == "合计", f"{ws.title}: 末行为合计行 (R{last})")
    check(ws.cell(last, 11).value == f"=SUM(K3:K{last-1})",
          f"{ws.title}: 合计公式 {ws.cell(last,11).value!r}")
    check(f"A{last}:H{last}" in [str(m) for m in ws.merged_cells.ranges],
          f"{ws.title}: A{last}:H{last} 已合并")

    hh = {ws.cell(2, i).value: i for i in range(1, 12)}
    n = 0; idxs = []; skeys = []
    for r in range(3, last):
        g = lambda k: ws.cell(r, hh[k]).value
        name = T(g("名称")); sp = T(g("规格型号")); mat = T(g("材质/牌号"))
        par = T(g("技术参数")); unit = T(g("单位")); qty = g("数量")
        note = T(g("备注"))
        price = g("单价"); tot_c = g("合价")
        if not (name or sp or mat or par): continue
        n += 1; idxs.append(g("序号"))
        q = float(qty); q = int(q) if q == int(q) else q
        k, mk = ident(name, sp + par, mat, unit, note)     # 注意：用「规格+参数」拼接再归一
        out_coarse[k].append((mk, q))
        out_members[k].append(dict(r=r, cat=g("类别"), name=name, spec=sp, par=par,
                                   mat=mat, note=note, unit=unit, qty=q))
        n_out += 1
        skeys.append((NAME_RANK.get((cat_name, name), 10 ** 6), name, dim(sp), mat, par, note, sp))
        if g("类别") != ws.title: check(False, f"{ws.title} R{r}: 类别列={g('类别')!r} 与表名不符")
        if price not in (None, ""): check(False, f"{ws.title} R{r}: 单价应为空，实为 {price!r}")
        if tot_c != f'=IF(J{r}="","",H{r}*J{r})': check(False, f"{ws.title} R{r}: 合价公式={tot_c!r}")
    check(idxs == list(range(1, n + 1)), f"{ws.title}: 序号 1..{n} 连续")
    check(skeys == sorted(skeys), f"{ws.title}: 行序符合「品类逻辑序 + 规格口径升序」")
    print(f"    数据行 {n}")

out_final = {}
dup_bucket = []
for k, pairs in out_coarse.items():
    ab = absorb(pairs)
    if len(ab) != 1:
        dup_bucket.append((k, pairs))
    for m, q in ab.items():
        out_final[(k, m)] = q

# ---------- 全局比对 ----------
print("\n=== 全局比对（身份 + 数量）===")
check(n_out == len(src_final), f"输出行数 {n_out} == 源合并后身份数 {len(src_final)}"
                               f"（源 {n_src} 行）")
check(not dup_bucket, f"每个分组键在输出中只对应 1 行（异常 {len(dup_bucket)} 个）")
for k, p in dup_bucket[:5]: print(f"       异常分组 {str(k)[:80]!r} {p}")

exp = Counter(src_final); act = Counter(out_final)
miss = exp - act; extra = act - exp
check(not miss, f"源侧身份全部出现在输出（缺失 {sum(miss.values())}）")
check(not extra, f"输出身份全部来自源表（多出 {sum(extra.values())}）")
for k, v in list(miss.items())[:8]: print(f"       缺失 {str(k)[:100]!r} x{v}")
for k, v in list(extra.items())[:8]: print(f"       多出 {str(k)[:100]!r} x{v}")

out_total = qsum([x["qty"] for ms in out_members.values() for x in ms])
check(out_total == src_total, f"数量总和守恒 源 {src_total} == 输出 {out_total}")

# 逐组核对数量 + 代表值无捏造
bad_qty, bad_rep, miss_row = 0, 0, 0
for (k, m), want in src_final.items():
    got = out_final.get((k, m))
    if got != want:
        bad_qty += 1
        print(f"       数量不符 {str(k)[:70]!r} 材质={m!r} 源={want} 输出={got}")
    mem = out_members.get(k, [])
    srcs = src_members.get(k, [])
    vals = {("n", T(x["name"])) for x in srcs} | {("s", T(x["spec"])) for x in srcs} \
           | {("m", T(x["mat"])) for x in srcs} | {("o", T(x["note"])) for x in srcs}
    for x in mem:
        for tag, v in (("n", x["name"]), ("m", x["mat"]), ("o", x["note"])):
            if v and (tag, v) not in vals:
                bad_rep += 1
                print(f"       代表值疑似捏造 {tag}={v!r} 不在源组内: {str(k)[:60]!r}")
check(bad_qty == 0, f"{len(src_final)} 个身份的数量均为源同身份数量之和（不符 {bad_qty} 组）")
check(bad_rep == 0, f"输出行的名称/材质/备注均取自源成员原文（捏造 {bad_rep} 处）")

# 源每行都被覆盖
uncovered = 0
for k, ms in src_members.items():
    for x in ms:
        if k not in out_members:
            uncovered += 1
check(uncovered == 0, f"源表 {n_src} 行逐行回溯均有归属（未覆盖 {uncovered} 行）")

print(f"\n分类分布: {dict(Counter(m['cat'] for ms in out_members.values() for m in ms))}")
print(f"\n{'>>> 全部校验通过' if not FAIL else '>>> 存在 %d 项失败' % len(FAIL)}")
for m in FAIL: print("   -", m)
sys.exit(1 if FAIL else 0)
