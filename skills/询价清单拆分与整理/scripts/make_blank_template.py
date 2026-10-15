# -*- coding: utf-8 -*-
"""从xlsx格式参考创建无业务数据的10列供应商空白模板。
用法：python make_blank_template.py 参考.xlsx 输出.xlsx
xls参考须先用可用的Excel/转换器另存临时副本，源文件保持不变。
"""
from pathlib import Path
from copy import copy
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HEADERS = ["序号", "材料名称", "规格型号", "材质/牌号", "技术参数", "单位", "数量", "备注", "单价", "金额"]
WIDTHS = [6, 29, 28, 20, 55, 8, 12, 35, 13, 16]

def create(reference, output):
    assert reference.resolve() != output.resolve(), "模板源文件不得覆盖"
    source = load_workbook(reference).active
    old = {c.value: c for c in source[2] if c.value}
    aliases = {"材料名称": "名称", "金额": "合价"}
    wb = Workbook(); ws = wb.active; ws.title = "类别名"
    ws.merge_cells("A1:J1")
    ws["A1"] = "项目全称\n（含13%增值税含运费，货到现汇付款，地址：交货地址）"
    ws["A1"].font = Font(name="宋体", size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 72
    line = Side(style="thin", color="BAC5CE")
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(2, col, name)
        template = old.get(name) or old.get(aliases.get(name))
        if template is not None:
            cell.border = copy(template.border)
        cell.font = Font(name="宋体", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF2F5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in [2, 3, 4]:
            ws.cell(row, col).border = Border(left=line, right=line, top=line, bottom=line)
            if row != 2:
                ws.cell(row, col).font = Font(name="宋体", size=10)
        ws.column_dimensions[cell.column_letter].width = WIDTHS[col - 1]
    ws["I3"].fill = PatternFill("solid", fgColor="E8F2FA")
    ws["I3"].font = Font(name="宋体", size=10, color="0000CC")
    ws["J3"] = '=IF(OR(I3="",G3="",NOT(ISNUMBER(G3))),"",G3*I3)'
    ws.merge_cells("A4:I4"); ws["A4"] = "合计"
    ws["J4"] = '=IF(COUNT(J3:J3)=0,"",SUM(J3:J3))'
    for address in ["I3", "J3", "J4"]:
        ws[address].number_format = "#,##0.00"
    ws.row_dimensions[2].height = 29; ws.row_dimensions[3].height = 33
    ws.freeze_panes = "C3"; ws.auto_filter.ref = "A2:J3"; ws.print_title_rows = "1:2"
    wb.save(output)
    check = load_workbook(output).active
    assert [c.value for c in check[2]] == HEADERS
    assert check.max_column == 10 and check["I3"].value is None
    assert check["G3"].value is None and check["B3"].value is None
    assert check["J3"].value == ws["J3"].value
    assert not any(c.comment for row in check for c in row)
    print("Created supplier template:", output)

if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    create(Path(sys.argv[1]), Path(sys.argv[2]))
