# -*- coding: utf-8 -*-
"""精确核查：按表头名定位列，按名称匹配行，检查源规格列全文是否都在新表三列中。"""
import sys, re
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

s = load_workbook(r"D:\UserData\Desktop\20261005中煤鄂尔多斯火炬气治理改造项目.xlsx", data_only=True)
o = load_workbook(r"D:\UserData\Desktop\20261005中煤鄂尔多斯火炬气治理改造项目_整理版.xlsx", data_only=True)

def c(ws, r, col):
    v = ws.cell(r, col).value
    return '' if v is None else str(v).replace('\r\n', '\n').strip()

def hdr(wo):
    return {c(wo, 1, i): i for i in range(1, wo.max_column + 1)}

# 源表规格列定位
SPEC = {
    '工程量': (range(6, 40), lambda ws, r: c(ws, r, 3)),
    '配电柜': (range(2, 18), lambda ws, r: c(ws, r, 4) if r <= 10 else c(ws, r, 5)),
    '线缆':   (range(1, 39), lambda ws, r: c(ws, r, 5)),
    '钢材':   (range(3, 59), lambda ws, r: c(ws, r, 6) + ' ' + (c(ws, r, 7) if r > 27 else '')),
    '灯具':   (range(4, 9),  lambda ws, r: c(ws, r, 6)),
    '阀门':   (range(4, 15), lambda ws, r: c(ws, r, 6)),
    '桥架':   (range(4, 7),  lambda ws, r: c(ws, r, 5)),
    '其他':   (range(2, 24), lambda ws, r: c(ws, r, 6)),
    '管件1':  (range(2, 89), lambda ws, r: c(ws, r, 4) + ' ' + c(ws, r, 5)),
    '保温':   (range(2, 8),  lambda ws, r: c(ws, r, 5) + ' ' + c(ws, r, 6)),
    '仪表':   (range(9, 21), lambda ws, r: c(ws, r, 3)),
}
NAME_COL = {k: 2 for k in SPEC}
NAME_COL.update({'工程量': 2, '配电柜': None, '线缆': 4, '钢材': 5, '灯具': 5,
                 '阀门': 5, '桥架': 4, '其他': 5, '管件1': 3, '保温': 4, '仪表': 2})

total_checked = 0
total_loss = 0
for sh, (rng, getraw) in SPEC.items():
    ws = s[sh]
    wo = o[sh]
    h = hdr(wo)
    iS, iM, iP, iN = h['规格型号'], h['材质/牌号'], h['技术参数'], h['名称']
    # 新表按名称建索引
    idx = {}
    for r in range(2, wo.max_row + 1):
        idx.setdefault(c(wo, r, iN), []).append(r)
    used = {}
    losses = 0
    for r in rng:
        raw = getraw(ws, r)
        if not raw:
            continue
        nm = c(ws, r, 3) if sh == '配电柜' and r <= 10 else (
             c(ws, r, 4) if sh == '配电柜' else c(ws, r, NAME_COL[sh]))
        if nm not in idx:
            print(f"  !! {sh}!R{r} 名称 '{nm[:16]}' 在新表找不到")
            continue
        k = used.get(nm, 0)
        used[nm] = k + 1
        orow = idx[nm][k] if k < len(idx[nm]) else idx[nm][-1]
        joined = re.sub(r'\s+', '', c(wo, orow, iS) + c(wo, orow, iM) + c(wo, orow, iP))
        total_checked += 1
        for tok in re.findall(r'[\u4e00-\u9fa5]{3,}|[A-Za-z0-9\.\-]{4,}', raw):
            if re.sub(r'\s+', '', tok) not in joined:
                print(f"  丢失 {sh}!R{r} '{nm[:14]}' token='{tok}'")
                print(f"       新规格={c(wo,orow,iS)[:50]!r}")
                print(f"       新材质={c(wo,orow,iM)[:50]!r}")
                print(f"       新参数={c(wo,orow,iP)[:50]!r}")
                losses += 1
                total_loss += 1
                break
    print(f"{sh:<8} 核对完成，丢失 {losses} 处")

print()
print(f"合计核对 {total_checked} 行，丢失 {total_loss} 处")
