# -*- coding: utf-8 -*-
"""
在合并版基础上：加单价/合价（含公式）+ 顶部大标题（仿参考文档格式）。
参考：D:\\工作\\材料\\询价\\已询价\\通州\\存的报价\\江苏通州湾港区热电联产项目（减压器）青岛志信，减温减压装置.xlsx
"""
import sys, re
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as gl

SRC = r"D:\UserData\Desktop\20261005中煤鄂尔多斯火炬气治理改造项目_合并版.xlsx"
OUT = r"D:\UserData\Desktop\中煤鄂尔多斯火炬气治理改造项目_材料计划（含13%增值税）.xlsx"

ORDER = ['设备', '工程量', '配电柜', '线缆', '钢材', '灯具', '阀门', '桥架',
         '其他', '管件1', '保温', '储罐', '仪表', '风机', '泵']

# 标题：项目名与地址按用户要求替换
TITLE = ('中煤鄂尔多斯能源化工有限公司乌审召化工分公司火炬气治理改造项目\n'
         '材料计划\n'
         '（含13%增值税含运费，货到付款，地址：内蒙古自治区鄂尔多斯市乌审旗乌审召镇'
         '中煤鄂尔多斯能源化工有限公司）')

YAHEI = 'Microsoft YaHei'
TITLE_FONT = Font(name=YAHEI, size=16, bold=True, color='FF000000')
HDR_FONT = Font(name=YAHEI, size=11, bold=True, color='FF000000')
HDR_FONT10 = Font(name=YAHEI, size=10, bold=True, color='FF000000')
BODY_FONT = Font(name=YAHEI, size=10)
FLAG_FONT = Font(name=YAHEI, size=10, color='C00000')
PRICE_FONT = Font(name=YAHEI, size=10, color='0000CC')      # 待填单价 蓝色
TOTAL_FONT = Font(name=YAHEI, size=10, color='000000')
THIN = Side(style='thin', color='FF000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NO_BORDER = Border(top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)
LEFT_CTR = Alignment(horizontal='left', vertical='center', wrap_text=True)
FLAG_FILL = PatternFill('solid', fgColor='FFF2CC')
MERGE_FILL = PatternFill('solid', fgColor='E2EFDA')
TITLE_FILL = PatternFill('solid', fgColor='FFFFFF')
PRICE_FILL = PatternFill('solid', fgColor='F2F7FF')          # 单价列淡蓝

def c(ws, r, col):
    v = ws.cell(r, col).value
    return '' if v is None else str(v).strip()

def is_num_or_blank(s):
    return (s == '') or bool(re.fullmatch(r'\d+(\.\d+)?', str(s).replace(',', '').strip()))

src = load_workbook(SRC, data_only=True)
out = Workbook()
out.remove(out.active)

report = []
for sh in ORDER:
    ws = src[sh]
    h = {ws.cell(1, i).value: i for i in range(1, ws.max_column + 1)}
    cols = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
    n = ws.max_row - 1

    # 新列顺序：行号 前置为序号，后面追加 单价、合价
    new_cols = ['序号'] + [x for x in cols if x != '行号'] + ['单价', '合价']
    # 找到关键列在 new_cols 中的位置
    ic = {cn: new_cols.index(cn) + 1 for cn in new_cols}
    i_qty = ic['数量']
    i_price = ic['单价']
    i_total = ic['合价']
    last_col = len(new_cols)

    ws2 = out.create_sheet(sh)
    # ---- 大标题：第1 行，A:末列 合并 ----
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    t = ws2.cell(1, 1)
    t.value = TITLE
    t.font = TITLE_FONT
    t.alignment = CENTER
    t.fill = TITLE_FILL
    for cidx in range(1, last_col + 1):
        cc = ws2.cell(1, cidx)
        cc.border = NO_BORDER
        cc.fill = TITLE_FILL
    ws2.cell(1, 1).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws2.cell(1, last_col).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws2.row_dimensions[1].height = 78

    # ---- 表头：第 2 行 ----
    for j, cn in enumerate(new_cols, 1):
        cell = ws2.cell(2, j)
        cell.value = cn
        cell.font = HDR_FONT10 if cn in ('单价', '合价') else HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    # ---- 数据 ----
    qty_bad = 0
    for i in range(1, n + 1):
        r = i + 1
        rowno = i + 2
        vals = []
        for cn in new_cols:
            if cn == '序号':
                vals.append(str(i))
            elif cn == '单价':
                vals.append(None)
            elif cn == '合价':
                q = c(ws, r, h['数量'])
                # 只有数量是纯数字才写公式；否则留空（数量是 "2+1" 这类复合值）
                if is_num_or_blank(q) and q != '':
                    qc = gl(i_qty)
                    pc = gl(i_price)
                    vals.append(f'=IF({pc}{rowno}="","",{qc}{rowno}*{pc}{rowno})')
                else:
                    vals.append(None)
                    if q != '':
                        qty_bad += 1
            else:
                vals.append(c(ws, r, h[cn]))
        ws2.append(vals)

        # 样式
        orig_fill = ws.cell(r, 1).fill
        is_merge_row = False
        try:
            is_merge_row = (orig_fill.fgColor.rgb or '').endswith('E2EFDA')
        except Exception:
            pass
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(rowno, j)
            cell.border = BORDER
            cn = new_cols[j - 1]
            if cn == '单价':
                cell.font = PRICE_FONT
                cell.fill = PRICE_FILL
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif cn == '合价':
                cell.font = TOTAL_FONT
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif cn == '待补标记':
                cell.font = FLAG_FONT
                cell.alignment = LEFT_TOP
                if v:
                    cell.fill = FLAG_FILL
            elif cn in ('技术参数', '备注', '规格型号', '材质/牌号', '类别', '设备位号', '名称'):
                cell.font = BODY_FONT
                cell.alignment = LEFT_TOP
                if is_merge_row:
                    cell.fill = MERGE_FILL
            else:
                cell.font = BODY_FONT
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if is_merge_row:
                    cell.fill = MERGE_FILL

    # ---- 合计行 ----
    tr = n + 3
    ws2.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=i_qty)
    cell = ws2.cell(tr, 1)
    cell.value = '合计'
    cell.font = Font(name=YAHEI, size=10, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    # 先写公式与样式，最后再统一描边（避免遍历合并单元格时抛异常）
    tc = ws2.cell(tr, i_total)
    tc.value = f'=SUM({gl(i_total)}3:{gl(i_total)}{n+2})'
    tc.font = Font(name=YAHEI, size=10, bold=True)
    tc.number_format = '#,##0.00'
    tc.alignment = Alignment(horizontal='right', vertical='center')
    for j in range(1, last_col + 1):
        cc = ws2.cell(tr, j)
        cc.border = BORDER
    ws2.row_dimensions[tr].height = 22

    # ---- 列宽 / 冻结 / 筛选 ----
    widths = {'序号': 6, '行号': 6, '设备位号': 15, '类别': 11, '名称': 26, '规格型号': 26,
              '材质/牌号': 22, '技术参数': 50, '单位': 7, '数量': 9, '备注': 16,
              '待补标记': 20, '单价': 12, '合价': 14}
    for j, cn in enumerate(new_cols, 1):
        ws2.column_dimensions[gl(j)].width = widths.get(cn, 12)
    ws2.freeze_panes = 'A3'
    ws2.auto_filter.ref = f"A2:{gl(last_col)}{n+2}"
    ws2.row_dimensions[2].height = 30
    report.append((sh, n, qty_bad))

out.save(OUT)

print(f"{'工作表':<10}{'数据行':>7}{'合价留空(数量非数字)':>22}")
tn = tb = 0
for sh, n, b in report:
    print(f'{sh:<10}{n:>7}{b:>22}')
    tn += n
    tb += b
print(f"{'合计':<10}{tn:>7}{tb:>22}")
print()
print('列顺序:', new_cols)
print('已保存:', OUT)
