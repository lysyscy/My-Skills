# -*- coding: utf-8 -*-
"""校验：大标题、单价/合价列、公式、合计、数据完整性。"""
import sys, re
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as gl

SRC = r"D:\UserData\Desktop\20261005中煤鄂尔多斯火炬气治理改造项目_合并版.xlsx"
OUT = r"D:\UserData\Desktop\中煤鄂尔多斯火炬气治理改造项目_材料计划（含13%增值税）.xlsx"

a = load_workbook(SRC, data_only=True)
b = load_workbook(OUT)          # 读公式
ORDER = ['设备', '工程量', '配电柜', '线缆', '钢材', '灯具', '阀门', '桥架',
         '其他', '管件1', '保温', '储罐', '仪表', '风机', '泵']
fail = 0

def c(ws, r, col):
    v = ws.cell(r, col).value
    return '' if v is None else str(v).strip()

print("=== 1. 大标题 ===")
ws = b['工程量']
h = {ws.cell(2, i).value: i for i in range(1, ws.max_column + 1)}
t = ws['A1']
lastcol = gl(ws.max_column)
print(f'  A1 内容: {str(t.value)[:100]!r}')
print(f'  合并范围: A1:{lastcol}1  实际: {[str(r) for r in ws.merged_cells.ranges if r.min_row==1]}')
print(f'  字体: {t.font.name} {t.font.sz} bold={t.font.b}')
print(f'  对齐: {t.alignment.horizontal}/{t.alignment.vertical}/wrap={t.alignment.wrap_text}')
print(f'  行高: {ws.row_dimensions[1].height}')
ck = ['中煤鄂尔多斯', '内蒙古自治区鄂尔多斯市乌审旗乌审召镇', '中煤鄂尔多斯能源化工有限公司',
      '材料计划', '含13%增值税含运费', '货到付款']
for k in ck:
    ok = k in str(t.value)
    if not ok: fail += 1
    print(f'  含「{k[:20]}」{"OK" if ok else "❌"}')
bad = [k for k in ('通州湾', '江苏省', '通州') if k in str(t.value)]
if bad: fail += 1
print(f'  旧项目名/地址已清除 {"OK" if not bad else "❌仍含 "+str(bad)}')
# 所有表都要有大标题
for sh in ORDER:
    w = b[sh]
    if not w['A1'].value or '火炬气治理改造' not in str(w['A1'].value):
        print(f'  ❌ {sh} 缺大标题')
        fail += 1
    if not any(r.min_row == 1 and r.max_col >= 10 for r in w.merged_cells.ranges):
        print(f'  ❌ {sh} 标题未跨列合并')
        fail += 1
print('  15 张表标题与合并：', 'OK' if fail == 0 else '有缺失')

print()
print("=== 2. 列顺序与表头 ===")
print(f'  {list(h.keys())}')
need = ['单价', '合价']
forbidden = ['报价供应商', '来源', '原序号', '原行号']
for k in forbidden:
    if k in h:
        fail += 1
        print(f'  ❌ 存在不该有的列: {k}')
print('  多余列检查：', 'OK' if not any(k in h for k in forbidden) else '有多余列')
for k in need:
    if k not in h: fail += 1; print(f'  ❌ 缺列 {k}')
print('  新列：OK' if all(k in h for k in need) else '  新列：有缺失')

print()
print("=== 3. 合价公式 ===")
tot_f = 0
for sh in ORDER:
    w = b[sh]
    hh = {w.cell(2, i).value: i for i in range(1, w.max_column + 1)}
    iq, ip, it = hh['数量'], hh['单价'], hh['合价']
    n = w.max_row - 3      # 去掉标题+表头+合计行
    okc = 0
    for r in range(3, 3 + n):
        q = c(w, r, iq)
        f = w.cell(r, it).value
        if q == '':
            continue
        if f and str(f).startswith('=IF('):
            exp = f'=IF({gl(ip)}{r}="","",{gl(iq)}{r}*{gl(ip)}{r})'
            if str(f) == exp:
                okc += 1
            else:
                print(f'  ❌ {sh}!{gl(it)}{r} 公式错误: {f}')
                fail += 1
        else:
            print(f'  ❌ {sh}!{gl(it)}{r} 数量[{q}] 合价无公式: {f!r}')
            fail += 1
    tot_f += okc
    print(f'  {sh:<6} {okc}/{n} 行公式正确')
print(f'  合计 {tot_f} 行公式正确')

print()
print("=== 4. 合计行公式 ===")
for sh in ORDER:
    w = b[sh]
    hh = {w.cell(2, i).value: i for i in range(1, w.max_column + 1)}
    it = hh['合价']
    tr = w.max_row
    v = w.cell(tr, it).value
    n = w.max_row - 3
    exp = f'=SUM({gl(it)}3:{gl(it)}{n+2})'
    ok = str(v) == exp
    if not ok: fail += 1
    print(f'  {sh:<6} {str(v):<22} {"OK" if ok else "❌应为 "+exp}')

print()
print("=== 5. 数据完整性（对比合并版）===")
for sh in ORDER:
    wa = a[sh]
    wb2 = b[sh]
    ha = {wa.cell(1, i).value: i for i in range(1, wa.max_column + 1)}
    hb = {wb2.cell(2, i).value: i for i in range(1, wb2.max_column + 1)}
    bad = 0
    for cn in ('名称', '规格型号', '材质/牌号', '单位', '数量', '待补标记'):
        for i in range(1, wa.max_row):
            va = c(wa, i + 1, ha[cn])
            vb = c(wb2, i + 2, hb[cn])
            if va != vb:
                if bad < 3:
                    print(f'  ❌ {sh} R{i+1} {cn}: [{va[:18]}] vs [{vb[:18]}]')
                bad += 1
                fail += 1
    print(f'  {sh:<6} {wa.max_row-1} 行 x 6 字段  {"OK" if bad==0 else f"{bad} 处不符"}')

print()
print("=== 6. 数字格式与冻结 ===")
w = b['工程量']
hh = {w.cell(2, i).value: i for i in range(1, w.max_column + 1)}
print(f'  单价格式: {w.cell(3, hh["单价"]).number_format!r}  合价格式: {w.cell(3, hh["合价"]).number_format!r}')
print(f'  冻结: {w.freeze_panes}  筛选: {w.auto_filter.ref}')
print(f'  行数校验: 合并版{a["工程量"].max_row-1} vs 计划版{w.max_row-3}  '
      f'{"OK" if a["工程量"].max_row-1==w.max_row-3 else "❌"}')

print()
print("=== 校验结论 ===")
print("PASS" if fail == 0 else f"FAIL：{fail} 处问题")
